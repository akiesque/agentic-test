from pydantic import BaseModel, field_validator
from typing import Optional
import json
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

def parse_number(v):
    if v is None or v == "-":
        return None
    v = v.strip()
    is_negative = v.startswith("(") and v.endswith(")")  # accounting-style negatives
    v = v.strip("()").replace(",", "").replace("%", "")
    try:
        num = float(v)
        return -num if is_negative else num
    except ValueError:
        raise ValueError(f"Could not parse number from: {v}")

class LineItem(BaseModel):
    q2_2026: float
    q1_2026: float
    qoq_change: Optional[float] = None
    jun_2026_ytd: float
    jun_2025_ytd: float
    yoy_change: Optional[float] = None

    @field_validator("q2_2026", "q1_2026", "jun_2026_ytd", "jun_2025_ytd", "qoq_change", "yoy_change", mode="before")
    @classmethod
    def clean_number(cls, v):
        return parse_number(v)

    @field_validator("qoq_change", "yoy_change")
    @classmethod
    def reject_duplicated_values(cls, v, info):
        if v is not None and abs(v) > 500:
            raise ValueError(f"{info.field_name} looks like a raw figure, not a %: {v}")
        return v
# --- end replacement ---

def export_to_excel(validated_items, filename="globe_q2_2026.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Summary"

    headers = ["Results of Operations (Php Mn)", "Q2 2026", "Q1 2026", "QoQ Change", "30-Jun 2026", "30-Jun 2025", "YoY Change"]
    ws.append(headers)

    for label, item in validated_items:
        ws.append([
            label,
            item.q2_2026,
            item.q1_2026,
            item.qoq_change if item.qoq_change is not None else "-",
            item.jun_2026_ytd,
            item.jun_2025_ytd,
            item.yoy_change if item.yoy_change is not None else "-",
        ])

    chart = BarChart()
    chart.title = "Q2 2026 vs Q1 2026"
    chart.y_axis.title = "Php Mn"
    chart.x_axis.title = "Line Item"

    data_ref = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=ws.max_row)
    categories = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(categories)

    ws.add_chart(chart, "I2")

    wb.save(filename)
    print(f"Saved to {filename}")

raw_json = '''
[
    {"Operating Revenues": {"Q2 2026": "46,955", "Q1 2026": "45,706", "QoQ Change": "3%", "30-Jun 2026": "92,661", "30-Jun 2025": "87,226", "YoY Change": "6%"}},
    {"Service Revenues": {"Q2 2026": "43,402", "Q1 2026": "41,965", "QoQ Change": "3%", "30-Jun 2026": "85,367", "30-Jun 2025": "80,188", "YoY Change": "6%"}},
    {"Mobile*": {"Q2 2026": "30,428", "Q1 2026": "29,965", "QoQ Change": "2%", "30-Jun 2026": "60,393", "30-Jun 2025": "57,073", "YoY Change": "6%"}},
    {"Home Broadband**": {"Q2 2026": "6,265", "Q1 2026": "6,178", "QoQ Change": "1%", "30-Jun 2026": "12,443", "30-Jun 2025": "11,708", "YoY Change": "6%"}},
    {"Corporate Data": {"Q2 2026": "5,910", "Q1 2026": "5,135", "QoQ Change": "15%", "30-Jun 2026": "11,045", "30-Jun 2025": "9,627", "YoY Change": "15%"}},
    {"Fixed line Voice": {"Q2 2026": "301", "Q1 2026": "322", "QoQ Change": "-7%", "30-Jun 2026": "623", "30-Jun 2025": "626", "YoY Change": "-1%"}},
    {"Others***": {"Q2 2026": "498", "Q1 2026": "365", "QoQ Change": "37%", "30-Jun 2026": "863", "30-Jun 2025": "1,154", "YoY Change": "-25%"}},
    {"Non-Service Revenues": {"Q2 2026": "3,553", "Q1 2026": "3,741", "QoQ Change": "-5%", "30-Jun 2026": "7,294", "30-Jun 2025": "7,038", "YoY Change": "4%"}},
    {"Costs and Expenses": {"Q2 2026": "24,254", "Q1 2026": "23,532", "QoQ Change": "3%", "30-Jun 2026": "47,786", "30-Jun 2025": "45,086", "YoY Change": "6%"}},
    {"Cost of Sales": {"Q2 2026": "3,742", "Q1 2026": "3,858", "QoQ Change": "-3%", "30-Jun 2026": "7,600", "30-Jun 2025": "7,119", "YoY Change": "7%"}},
    {"Operating Expenses": {"Q2 2026": "20,512", "Q1 2026": "19,674", "QoQ Change": "4%", "30-Jun 2026": "40,186", "30-Jun 2025": "37,967", "YoY Change": "6%"}},
    {"EBITDA": {"Q2 2026": "22,701", "Q1 2026": "22,174", "QoQ Change": "2%", "30-Jun 2026": "44,875", "30-Jun 2025": "42,140", "YoY Change": "6%"}},
    {"EBITDA Margin": {"Q2 2026": "52.3%", "Q1 2026": "52.8%", "QoQ Change": "-", "30-Jun 2026": "52.6%", "30-Jun 2025": "52.6%", "YoY Change": "-"}},
    {"Depreciation": {"Q2 2026": "14,304", "Q1 2026": "14,361", "QoQ Change": "-", "30-Jun 2026": "28,665", "30-Jun 2025": "26,430", "YoY Change": "8%"}},
    {"EBIT": {"Q2 2026": "8,397", "Q1 2026": "7,813", "QoQ Change": "7%", "30-Jun 2026": "16,210", "30-Jun 2025": "15,710", "YoY Change": "3%"}},
    {"EBIT Margin": {"Q2 2026": "19.3%", "Q1 2026": "18.6%", "QoQ Change": "-", "30-Jun 2026": "19.0%", "30-Jun 2025": "19.6%", "YoY Change": "-"}},
    {"Non-Operating Income (Charges)": {"Q2 2026": "(1,783)", "Q1 2026": "(1,286)", "QoQ Change": "39%", "30-Jun 2026": "(3,069)", "30-Jun 2025": "(980)", "YoY Change": "213%"}},
    {"Net Income After Tax (NIAT)": {"Q2 2026": "5,485", "Q1 2026": "5,554", "QoQ Change": "-1%", "30-Jun 2026": "11,039", "30-Jun 2025": "12,437", "YoY Change":"-11%"}},
    {"Core Net Income": {"Q2 2026": "5,256", "Q1 2026": "4,932", "QoQ Change": "7%", "30-Jun 2026": "10,188", "30-Jun 2025": "10,431", "YoY Change": "-2%"}}
]
'''

data = json.loads(raw_json)
key_map = {
    "Q2 2026": "q2_2026", "Q1 2026": "q1_2026", "QoQ Change": "qoq_change",
    "30-Jun 2026": "jun_2026_ytd", "30-Jun 2025": "jun_2025_ytd", "YoY Change": "yoy_change",
}

validated_items = []
for item in data:
    label, values = list(item.items())[0]
    renamed = {key_map[k]: v for k, v in values.items() if k in key_map}
    try:
        validated = LineItem(**renamed)
        validated_items.append((label, validated))
    except Exception as e:
        print(f"❌ {label}: SKIPPED, flagged — {e}")

export_to_excel(validated_items)